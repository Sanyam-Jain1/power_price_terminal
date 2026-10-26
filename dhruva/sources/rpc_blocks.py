"""Block-level RPC data: every entity, every 15-minute block (and NERPC's 5-minute SRAS).

The weekly PDF accounts (rpc.py) are the settlement of record. Each RPC also
publishes the block-wise data the account was computed from; this module reads
those files and reconciles them back to the account:

  * per entity and day, against the day's "Total" row where the file has one (NERPC);
  * per entity and week, against the entity's payable/receivable in the weekly PDF.

An entity is matched to its PDF row by name, or failing that by amount (its
block sums equal exactly one PDF row's payable and receivable). Anything that
matches neither way is kept but flagged "unmatched", never re-labelled.

Units: energy MWh, money Rs, rates Rs/MWh.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import datetime

DSM_COLUMNS = ("date", "block", "entity", "actual_mwh", "schedule_mwh", "sras_mwh", "deviation_mwh",
               "payable_rs", "receivable_rs", "normal_rate_rs_mwh", "nuclear_adj_payable_rs", "nuclear_adj_receivable_rs")
SRAS_COLUMNS = ("date", "block", "provider", "up_mwh", "down_mwh", "net_mwh", "energy_rs")
SRAS5_COLUMNS = ("date", "slot", "provider", "up_mwh", "down_mwh", "mileage_mwh", "incentive_rs")
TRAS_COLUMNS = ("date", "block", "provider",
                "da_up_cleared_mwh", "da_up_sched_mwh", "da_down_mwh", "da_up_rs_mwh", "da_down_rs_mwh", "da_rs",
                "rt_up_cleared_mwh", "rt_up_sched_mwh", "rt_down_mwh", "rt_up_rs_mwh", "rt_down_rs_mwh", "rt_rs",
                "sf_up_mwh", "sf_down_mwh", "sf_rs", "em_up_mwh", "em_down_mwh", "em_rs")
SCUC_COLUMNS = ("date", "block", "generator", "increment_mwh", "decrement_mwh", "variable_cost_rs_mwh", "net_rs")


def _f(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _iso(v) -> str | None:
    s = str(v or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _col(hdr: list[str], *pats: str) -> int | None:
    """Index of the first header cell matching any pattern (case-insensitive, whitespace-collapsed)."""
    cells = [" ".join(str(h or "").replace("\\n", " ").split()).lower() for h in hdr]
    for p in pats:
        for i, c in enumerate(cells):
            if re.search(p, c):
                return i
    return None


def _xlsx(data: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


class Result(dict):
    """{status, message, rows, freq: {(date, block): [hz...]}, day_checks: [...]}"""


def _result(rows, freq, day_checks=None, skipped=None) -> Result:
    msg = []
    if skipped:
        msg.append(f"skipped {len(skipped)} file(s)/sheet(s) without block data: {', '.join(sorted(skipped)[:6])}"
                   + ("…" if len(skipped) > 6 else ""))
    return Result(status="ok" if rows else "parse_error", message="; ".join(msg) if rows else "no block rows found",
                  rows=rows, freq=freq, day_checks=day_checks or [])


def within(res: Result, start, end) -> Result:
    """Keep only rows (and checks) dated inside [start, end]; files can span more than one week."""
    a, b = start.isoformat(), end.isoformat()
    out = Result(res)
    out["rows"] = [r for r in res["rows"] if a <= r[0] <= b]
    out["freq"] = {k: v for k, v in res.get("freq", {}).items() if a <= k[0] <= b}
    out["day_checks"] = [c for c in res.get("day_checks", []) if a <= c["date"] <= b]
    if res["status"] == "ok" and not out["rows"]:
        out.update(status="parse_error", message=f"the data file has no rows for {a} to {b}")
    return out


# ------------------------------------------------------------------ DSM, per RPC layout

def erpc_dsm(data: bytes) -> Result:
    """ERPC block-wise workbook: one sheet per entity, Normal Rate in p/kWh."""
    rows, freq, skipped = [], defaultdict(list), []
    for ws in _xlsx(data).worksheets:
        c = None
        for r in ws.iter_rows(values_only=True):
            if c is None:
                if r and r[0] == "Date":
                    h = [str(x or "") for x in r]
                    c = {k: _col(h, p) for k, p in (("freq", r"^freq"), ("act", r"^actual"), ("sch", r"^schedule"),
                                                    ("sras", r"^sras"), ("dev", r"^deviation \(mwh"),
                                                    ("pay", r"dsm payable"), ("rec", r"dsm receivable"),
                                                    ("nr", r"^normal dsm rate"))}
                continue
            d = _iso(r[0]) if r else None
            if not d or not isinstance(r[2], (int, float)):
                continue
            g = lambda k: _f(r[c[k]]) if c[k] is not None and c[k] < len(r) else None  # noqa: E731
            b = int(r[2])
            nr = g("nr")
            rows.append((d, b, ws.title, g("act"), g("sch"), g("sras"), g("dev"), g("pay"), g("rec"),
                         None if nr is None else round(nr * 10, 2), None, None))
            if g("freq") is not None:
                freq[(d, b)].append(g("freq"))
        if c is None:
            skipped.append(ws.title)
    return _result(rows, freq, skipped=skipped)


def wrpc_dsm(data: bytes) -> Result:
    """WRPC zip of per-constituent CSVs (*_DSM-2024_Data.csv); Normal Rate in p/kWh."""
    z = zipfile.ZipFile(io.BytesIO(data))
    rows, freq, skipped = [], defaultdict(list), []
    data_files = [n for n in z.namelist() if n.endswith("_DSM-2024_Data.csv") and n.count("/") <= 1]
    # Usually the CSVs sit at the top level, and anything in a folder is a nested copy of an older
    # week; some weeks ship everything inside a single dated folder instead.
    top = [n for n in data_files if "/" not in n]
    for name in top or data_files:
        rd = csv.reader(io.StringIO(z.read(name).decode("utf-8", errors="replace")))
        hdr = next(rd, [])
        c = {k: _col(hdr, p) for k, p in (("date", r"^date$"), ("blk", r"^block$"), ("freq", r"^freq"),
                                          ("ent", r"^constituents?$"), ("act", r"^actual"), ("sch", r"^schedule"),
                                          ("sras", r"^sras"), ("dev", r"^deviation ?\(mwh"), ("pay", r"^dsm payable"),
                                          ("rec", r"^dsm receivable"), ("nr", r"^normal rate"),
                                          # States carry their share of the nuclear stations' DSM (settled via beneficiaries).
                                          ("apay", r"^adjusted dsm payable"), ("arec", r"^adjusted dsm receivable"))}
        if c["date"] is None or c["blk"] is None or c["pay"] is None:
            skipped.append(name)
            continue
        fallback = name.split("/")[-1].removesuffix("_DSM-2024_Data.csv")
        for r in rd:
            if len(r) <= c["pay"]:
                continue
            d = _iso(r[c["date"]])
            try:
                b = int(r[c["blk"]])
            except ValueError:
                continue
            g = lambda k: _f(r[c[k]]) if c[k] is not None and c[k] < len(r) else None  # noqa: E731
            nr = g("nr")
            ent = (r[c["ent"]].strip() if c["ent"] is not None else "") or fallback
            rows.append((d, b, ent, g("act"), g("sch"), g("sras"), g("dev"), g("pay"), g("rec"),
                         None if nr is None else round(nr * 10, 2), g("apay"), g("arec")))
            if g("freq") is not None:
                freq[(d, b)].append(g("freq"))
    return _result(rows, freq, skipped=skipped)


def nerpc_dsm(members: dict[str, bytes]) -> Result:
    """NERPC DeviationDetails workbook(s): one sheet per entity, a 'Total' row opening each day."""
    rows, freq, checks, skipped = [], defaultdict(list), [], []
    for n, data in sorted(members.items()):
        if not n.startswith("DeviationDetails"):
            continue
        for ws in _xlsx(data).worksheets:
            c, day, totals = None, None, []
            got = defaultdict(lambda: [0.0, 0.0, 0.0])
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                if r[0] == "Time":
                    h = [str(x or "") for x in r]
                    # Inter-regional sheets have no block column; their rows are keyed by start time.
                    c = {k: _col(h, p) for k, p in (("blk", r"^block number"), ("freq", r"^frequency \(hz"), ("act", r"^actual"),
                                                    ("sch", r"^scheduled"), ("dev", r"^total deviation"),
                                                    ("nr", r"^normal rate"), ("rec", r"^amount rec"),
                                                    ("pay", r"^amount pay"))}
                    continue
                if r[1] == "Total" and _iso(r[0]):
                    day = _iso(r[0])
                    totals.append((day, r))
                    continue
                if not (day and c):
                    continue
                if c["blk"] is not None:
                    if not isinstance(r[c["blk"]], (int, float)):
                        continue
                    b = int(r[c["blk"]])
                else:
                    m = re.fullmatch(r"(\d{2}):(\d{2})", str(r[0] or ""))
                    if not m:
                        continue
                    b = (int(m.group(1)) * 60 + int(m.group(2))) // 15 + 1
                g = lambda k: _f(r[c[k]]) if c[k] is not None and c[k] < len(r) else None  # noqa: E731
                rows.append((day, b, ws.title, g("act"), g("sch"), None, g("dev"), g("pay") or 0.0, g("rec") or 0.0, g("nr"), None, None))
                acc = got[day]
                acc[0] += g("dev") or 0
                acc[1] += g("pay") or 0
                acc[2] += g("rec") or 0
                if g("freq") is not None:
                    freq[(day, b)].append(g("freq"))
            if c is None:
                skipped.append(ws.title)
                continue
            for day, r in totals:
                t = lambda k: _f(r[c[k]]) if c[k] is not None and c[k] < len(r) else None  # noqa: E731
                want = (t("dev"), t("pay") or 0.0, t("rec") or 0.0)
                have = got[day]
                ok = all(w is None or abs(w - x) <= max(1.0, abs(w) * 1e-6) for w, x in zip(want, have))
                checks.append({"entity": ws.title, "date": day, "ok": ok,
                               "printed": {"deviation_mwh": want[0], "payable_rs": want[1], "receivable_rs": want[2]},
                               "summed": {"deviation_mwh": round(have[0], 3), "payable_rs": round(have[1], 2),
                                          "receivable_rs": round(have[2], 2)}})
    return _result(rows, freq, checks, skipped)


# ------------------------------------------------------------------ NERPC ancillary data files

def nerpc_sras(members: dict[str, bytes]) -> tuple[Result, Result]:
    """15-minute energy (SRAS Up/Down/Net, variable charges) and 5-minute mileage and incentive."""
    r15, r5 = [], []
    for n, data in sorted(members.items()):
        if not n.upper().startswith("SRAS"):
            continue
        for ws in _xlsx(data).worksheets:
            five = ws.title.lower().endswith("_5min")
            c = None
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                if r[0] == "Date":
                    h = [str(x or "") for x in r]
                    c = {k: _col(h, p) for k, p in (("ent", r"^entity"), ("blk", r"^block"), ("up", r"^sras up"),
                                                    ("dn", r"^sras down"), ("net", r"^sras net"), ("mil", r"mileage"),
                                                    ("rs", r"^variable charges|^incentive"))}
                    continue
                d = _iso(r[0])
                if not (c and d) or not isinstance(r[c["blk"]], (int, float)):
                    continue
                g = lambda k: _f(r[c[k]]) if c[k] is not None and c[k] < len(r) else None  # noqa: E731
                prov = str(r[c["ent"]] or ws.title.removesuffix("_5Min")).strip()
                if five:
                    r5.append((d, int(r[c["blk"]]), prov, g("up"), g("dn"), g("mil"), g("rs")))
                else:
                    r15.append((d, int(r[c["blk"]]), prov, g("up"), g("dn"), g("net"), g("rs")))
    return _result(r15, {}), _result(r5, {})


def nerpc_tras(members: dict[str, bytes]) -> Result:
    """Per provider per block: DA and RT AS-market TRAS, plus shortfall and emergency TRAS."""
    rows = []
    for n, data in sorted(members.items()):
        if not n.upper().startswith("TRAS"):
            continue
        for ws in _xlsx(data).worksheets:
            for r in ws.iter_rows(min_row=3, values_only=True):
                d = _iso(r[0]) if r else None
                if not d or not isinstance(r[1], (int, float)):
                    continue
                g = lambda i: _f(r[i]) if i < len(r) else None  # noqa: E731
                z = lambda i: g(i) or 0.0  # noqa: E731
                rows.append((d, int(r[1]), ws.title,
                             g(2), g(3), g(4), g(5) if z(2) > 0 else None, g(7) if z(4) > 0 else None, z(8) + z(9) - z(10),
                             g(11), g(12), g(13), g(14) if z(11) > 0 else None, g(16) if z(13) > 0 else None, z(17) + z(18) - z(19),
                             g(20), g(21), z(23) - z(24),
                             g(25), g(26), z(28) - z(29)))
    return _result(rows, {})


def nerpc_scuc(members: dict[str, bytes]) -> Result:
    rows = []
    for n, data in sorted(members.items()):
        if not n.upper().startswith("SCUC"):
            continue
        for ws in _xlsx(data).worksheets:
            for r in ws.iter_rows(values_only=True):
                d = _iso(r[0]) if r else None
                if not d or not isinstance(r[1], (int, float)):
                    continue
                g = lambda i: _f(r[i]) if i < len(r) else None  # noqa: E731
                vc = g(4)
                rows.append((d, int(r[1]), ws.title, g(2), g(3), None if vc is None else round(vc * 10, 2), g(7)))
    return _result(rows, {})


# ------------------------------------------------------------------ reconciliation with the weekly PDF

def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower().replace("&", "and"))


def _close(a: float, b: float, tol_rs: float) -> bool:
    return abs(a - b) <= max(tol_rs, abs(b) * 0.002)


def reconcile(rows: list[tuple], pdf_rows: list[dict], tol_rs: float = 1000.0) -> dict:
    """Match each block-data entity to its row in the weekly account and compare payable/receivable.

    tol_rs absorbs rounding in the PDF (ERPC prints Rs lakh to 2 decimals, i.e. Rs 1,000).
    Returns {entities: {block_name: {...}}, summary: {...}}.
    """
    sums: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0])
    for r in rows:
        acc = sums[r[2]]
        acc[0] += (r[7] or 0) + (r[10] or 0)
        acc[1] += (r[8] or 0) + (r[11] or 0)
        acc[2] += abs(r[6] or 0)
    pdf = [p for p in pdf_rows if p.get("name")]
    by_name = {_norm(p["name"]): p for p in pdf}
    used: set[str] = set()
    out: dict[str, dict] = {}
    # Pass 1: names. Pass 2: amounts, only where exactly one unused PDF row fits.
    for ent, (pay, rec, _) in sums.items():
        p = by_name.get(_norm(ent))
        if p and p["name"] not in used:
            used.add(p["name"])
            out[ent] = {"pdf_name": p["name"], "category": p.get("category"), "match": "name"}
    for ent, (pay, rec, _) in sums.items():
        if ent in out:
            continue
        fits = [p for p in pdf if p["name"] not in used and (pay or rec) and (
                (_close(pay, p.get("payable_rs") or 0, tol_rs) and _close(rec, p.get("receivable_rs") or 0, tol_rs))
                or _close(pay - rec, (p.get("payable_rs") or 0) - (p.get("receivable_rs") or 0), tol_rs))]
        if len(fits) == 1:
            used.add(fits[0]["name"])
            out[ent] = {"pdf_name": fits[0]["name"], "category": fits[0].get("category"), "match": "amount"}
        else:
            out[ent] = {"pdf_name": None, "category": None, "match": "unmatched"}
    pdf_by = {p["name"]: p for p in pdf}
    n_ok = n_bad = 0
    for ent, m in out.items():
        pay, rec, absdev = sums[ent]
        m.update(payable_rs=round(pay, 2), receivable_rs=round(rec, 2), abs_deviation_mwh=round(absdev, 3))
        if m["pdf_name"]:
            p = pdf_by[m["pdf_name"]]
            pp, pr = p.get("payable_rs") or 0, p.get("receivable_rs") or 0
            m["pdf_payable_rs"], m["pdf_receivable_rs"] = pp, pr
            # Some accounts (NERPC) print each entity's weekly net on one side only.
            gross = _close(pay, pp, tol_rs) and _close(rec, pr, tol_rs)
            net = (not pp or not pr) and _close(pay - rec, pp - pr, tol_rs)
            m["ties"] = gross or net
            m["basis"] = "gross" if gross else "net" if net else None
            n_ok += m["ties"]
            n_bad += not m["ties"]
        else:
            m["ties"] = None
    missing = sorted(set(pdf_by) - used)
    return {"entities": out, "summary": {"entities": len(out), "ties": n_ok, "differs": n_bad,
                                         "unmatched": sum(1 for m in out.values() if m["match"] == "unmatched"),
                                         "pdf_rows_without_block_data": missing}}
